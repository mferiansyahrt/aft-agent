"""Ingest Excel CP Prima ke SQLite.

Contoh:
    python3 source/database/ingest.py --input docs_revised --db source/database/aft.sqlite

Database dibuat ulang dari nol setiap kali dijalankan.
"""
import argparse
import hashlib
import sqlite3
import zipfile
import sys
from datetime import datetime
from pathlib import Path

import sheet_spec as spec
from extract import extract
from openpyxl.utils.exceptions import InvalidFileException

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]


def insert(cur, table, rec):
    cols = list(rec)
    cur.execute(f"INSERT INTO {table} ({', '.join(cols)}) VALUES ({', '.join('?' * len(cols))})",
                [rec[c] for c in cols])
    return cur.lastrowid


def check_schema(con):
    """Pastikan kolom tabel lebar di schema.sql sama dengan sheet_spec.py."""
    for ws in spec.WIDE_SHEETS:
        info = list(con.execute(f"PRAGMA table_info({ws['table']})"))
        db_cols = [r[1] for r in info if not r[5]]  # r[5] = flag primary key
        expected = [f for _, f, _ in ws["cols"].values()]
        extra = [c for c in db_cols if c not in expected + ["file_id", "source_row"]]
        missing = [c for c in expected if c not in db_cols]
        if missing or extra:
            sys.exit(f"schema.sql tidak sesuai sheet_spec untuk {ws['table']}: kurang={missing} lebih={extra}")


def ingest_file(cur, path):
    data = extract(path)
    file_id = insert(cur, "source_file", {
        "file_name": path.name, "file_kind": data["kind"],
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "ingested_at": datetime.now().isoformat(timespec="seconds"),
    })
    issues = list(data["issues"])
    counts = {}

    if data["kind"] == "WPA":
        # kode SAP bisa muncul lebih dari sekali di daftar BA-BB; simpan sekali, deskripsi harus sama
        known = dict(cur.execute("SELECT sap_code, description FROM material"))
        for m in data["materials"]:
            if m["sap_code"] not in known:
                insert(cur, "material", {**m, "file_id": file_id})
                known[m["sap_code"]] = m["description"]
            elif known[m["sap_code"]] != m["description"]:
                issues.append({"sheet_name": spec.WPA_REPORT["sheet"], "cell_ref": f"BB{m['source_row']}",
                               "raw_value": m["description"], "expected_type": "TEXT",
                               "reason": f"kode SAP {m['sap_code']} duplikat dengan deskripsi berbeda "
                                         f"({known[m['sap_code']]!r})"})
        counts["material"] = len(set(m["sap_code"] for m in data["materials"]))

        report_id = insert(cur, "wpa_report", {**data["report"], "file_id": file_id})
        for lot in data["lots"]:
            quality = lot.pop("quality")
            if lot["sap_code"] is not None and lot["sap_code"] not in known:
                issues.append({"sheet_name": spec.WPA_REPORT["sheet"], "cell_ref": f"C{lot['source_row']}",
                               "raw_value": str(lot["sap_code"]), "expected_type": "INTEGER",
                               "reason": "kode SAP tidak ada di daftar BA-BB"})
                lot["sap_code"] = None
            lot_id = insert(cur, "wpa_stock_lot", {**lot, "report_id": report_id, "file_id": file_id})
            cur.executemany("INSERT INTO wpa_stock_lot_quality (lot_id, parameter_code, value, source_col) VALUES (?,?,?,?)",
                            [(lot_id, p, v, c) for p, v, c in quality])
        counts["wpa_stock_lot"] = len(data["lots"])

        for rec in data["receipts"]:
            tests = rec.pop("tests")
            rid = insert(cur, "rm_receipt", {**rec, "file_id": file_id})
            cur.executemany("INSERT INTO rm_test_result (receipt_id, parameter_code, value, source_col) VALUES (?,?,?,?)",
                            [(rid, p, v, c) for p, v, c in tests])
        counts["rm_receipt"] = len(data["receipts"])
    else:
        for ws in spec.WIDE_SHEETS:
            for rec in data[ws["table"]]:
                insert(cur, ws["table"], {**rec, "file_id": file_id})
            counts[ws["table"]] = len(data[ws["table"]])

    for issue in issues:
        insert(cur, "ingest_issue", {**issue, "file_id": file_id})
    counts["ingest_issue"] = len(issues)
    return data["kind"], counts


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--input", default=str(ROOT / "docs_revised"), help="folder berisi file .xlsx")
    ap.add_argument("--db", default=str(HERE / "aft.sqlite"), help="path file SQLite")
    args = ap.parse_args()

    files = sorted(p for p in Path(args.input).glob("*.xlsx") if not p.name.startswith("~$"))
    if not files:
        sys.exit(f"tidak ada file .xlsx di {args.input}")

    db = Path(args.db)
    db.parent.mkdir(parents=True, exist_ok=True)
    db.unlink(missing_ok=True)
    con = sqlite3.connect(db)
    con.executescript((HERE / "schema.sql").read_text())
    check_schema(con)
    cur = con.cursor()
    cur.executemany("INSERT INTO parameter (code, name) VALUES (?, ?)", spec.PARAMETERS.items())

    for path in files:
        try:
            kind, counts = ingest_file(cur, path)
        except (zipfile.BadZipFile, InvalidFileException):
            con.close()
            sys.exit(f"GAGAL: '{path.name}' bukan file Excel yang valid ({path.stat().st_size} byte). "
                     f"Hapus atau ganti file ini di {args.input}.")
        except ValueError as exc:
            con.close()
            sys.exit(f"GAGAL: '{path.name}': {exc}. Folder input hanya boleh berisi Excel WPA dan Feed Performance.")
        print(f"{path.name} [{kind}]")
        for table, n in counts.items():
            print(f"  {table:16s} {n}")
    con.commit()
    con.close()
    print(f"selesai -> {db}")


if __name__ == "__main__":
    main()
