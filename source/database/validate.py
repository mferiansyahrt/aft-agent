"""Validasi isi SQLite terhadap Excel, sel per sel.

Contoh:
    python3 source/database/validate.py --input docs_revised --db source/database/aft.sqlite

Script ini sengaja tidak memakai extract.py: Excel dibaca ulang langsung dan setiap sel
dibandingkan dengan nilai di database lewat source_row/source_col. Yang dicek:
  1. teks header sesuai sheet_spec (tidak ada kolom tergeser)
  2. jumlah baris data Excel = jumlah baris DB
  3. setiap sel Excel = nilai DB, dan tidak ada nilai DB tanpa sel asal
  4. tidak ada sel terisi di kolom yang tidak ter-mapping
  5. tidak ada sheet tak dikenal dan tidak ada ingest_issue
Exit code 1 kalau ada satu saja temuan.
"""
import argparse
import math
import re
import sqlite3
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

import sheet_spec as spec

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
DATE_FORMATS = ("%d %B %Y", "%d %b %Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y")
MAX_SHOW = 15


def blank(v):
    return v is None or (isinstance(v, str) and not v.strip())


def header_text(v):
    return "" if v is None else re.sub(r"\s+", " ", str(v)).strip()


def iso(raw):
    if isinstance(raw, datetime):
        return raw.strftime("%Y-%m-%d") if raw.time() == datetime.min.time() else raw.isoformat(sep=" ")
    if isinstance(raw, date):
        return raw.strftime("%Y-%m-%d")
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(str(raw).strip(), fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return None


def same(raw, db, typ):
    """Apakah nilai DB sesuai dengan isi sel Excel untuk tipe kolom ini."""
    if blank(raw):
        return db is None
    if db is None:
        return False
    if typ in ("REAL", "INTEGER"):
        try:
            x = float(raw.strip()) if isinstance(raw, str) else float(raw)
        except (TypeError, ValueError):
            return False
        if typ == "INTEGER" and not isinstance(db, int):
            return False
        return isinstance(db, (int, float)) and math.isclose(x, db, rel_tol=1e-12, abs_tol=1e-12)
    if typ == "DATE":
        return db == iso(raw)
    # TEXT
    if not isinstance(db, str):
        return False
    if isinstance(raw, str):
        return db == raw.strip()
    if isinstance(raw, (datetime, date)):
        return db == iso(raw)
    if isinstance(raw, (int, float)):
        if float(raw).is_integer():
            return db == str(int(raw))
        return math.isclose(float(db), raw, rel_tol=1e-12)
    return db == str(raw)


class Sheet:
    def __init__(self, ws):
        self.name = ws.title
        self.rows = {i: r for i, r in enumerate(ws.iter_rows(values_only=True), start=1)}
        self.max_row = max(self.rows, default=0)

    def get(self, row, col):
        r = self.rows.get(row)
        i = column_index_from_string(col) - 1
        return r[i] if r is not None and i < len(r) else None

    def filled_cols(self, row):
        r = self.rows.get(row) or ()
        return [get_column_letter(i + 1) for i, v in enumerate(r) if not blank(v)]


class Result:
    def __init__(self, label):
        self.label = label
        self.excel_rows = self.db_rows = self.cells = 0
        self.diffs, self.unmapped, self.notes = [], [], []

    def diff(self, ref, msg, raw=None, db=None):
        self.diffs.append(f"{ref}: {msg}" + (f" | excel={raw!r} db={db!r}" if raw is not None or db is not None else ""))

    @property
    def ok(self):
        return not self.diffs and not self.unmapped


def check_header(res, sh, row, col, expected):
    got = header_text(sh.get(row, col))
    if got != expected:
        res.diff(f"{col}{row}", f"header beda, harap {expected!r}", got, None)


def check_coverage(res, sh, start, allowed, skip=None):
    for r in range(start, sh.max_row + 1):
        for col in sh.filled_cols(r):
            if col in allowed:
                continue
            if skip and col in skip:
                res.notes.append(f"{col}{r} (kolom {skip[col]}) tidak di-ingest")
            else:
                res.unmapped.append(f"{col}{r}={sh.get(r, col)!r}")


def check_values(res, sh, row, rec, cols):
    """cols: {huruf: (field, tipe)}; rec: dict kolom DB."""
    for col, (field, typ) in cols.items():
        raw = sh.get(row, col)
        res.cells += 1
        if not same(raw, rec[field], typ):
            res.diff(f"{col}{row}", field, raw, rec[field])


def check_long(res, sh, row, db_vals, cols):
    """db_vals: {source_col: value} untuk satu baris; cols: {huruf: kode parameter}."""
    seen = set()
    for col, param in cols.items():
        raw = sh.get(row, col)
        res.cells += 1
        got = db_vals.get(col)
        if blank(raw):
            if got is not None:
                res.diff(f"{col}{row}", f"{param} ada di DB padahal sel kosong", raw, got[1])
            continue
        if got is None:
            res.diff(f"{col}{row}", f"{param} tidak ada di DB", raw, None)
        elif got[0] != param:
            res.diff(f"{col}{row}", f"kode parameter salah ({got[0]} != {param})")
        elif not same(raw, got[1], "REAL"):
            res.diff(f"{col}{row}", param, raw, got[1])
        seen.add(col)
    for col in set(db_vals) - set(cols):
        res.diff(f"{col}{row}", "nilai DB dari kolom yang tidak ada di spec", None, db_vals[col])


def rows_by_source(con, sql, args):
    cur = con.execute(sql, args)
    names = [d[0] for d in cur.description]
    return {r[names.index("source_row")]: dict(zip(names, r)) for r in cur}


def compare_rows(res, excel_rows, db_rows):
    res.excel_rows, res.db_rows = len(excel_rows), len(db_rows)
    for r in sorted(set(excel_rows) - set(db_rows)):
        res.diff(f"baris {r}", "ada di Excel tapi tidak ada di DB")
    for r in sorted(set(db_rows) - set(excel_rows)):
        res.diff(f"baris {r}", "ada di DB tapi bukan baris data di Excel")
    return sorted(set(excel_rows) & set(db_rows))


# ---------------------------------------------------------------------- WPA
def validate_wpa(con, wb, file_id):
    out = []
    s = spec.WPA_REPORT
    sh = Sheet(wb[s["sheet"]])

    res = Result("WPA REPORT · metadata")
    for ref, label in s["meta_labels"].items():
        check_header(res, sh, int(re.sub(r"\D", "", ref)), re.sub(r"\d", "", ref), label)
    reports = list(con.execute("SELECT * FROM wpa_report WHERE file_id=?", (file_id,)))
    names = [d[1] for d in con.execute("PRAGMA table_info(wpa_report)")]
    res.excel_rows, res.db_rows = 1, len(reports)
    if len(reports) != 1:
        res.diff("wpa_report", f"harus 1 baris, ada {len(reports)}")
    else:
        rep = dict(zip(names, reports[0]))
        for field, (ref, _) in s["meta"].items():
            raw = sh.get(int(re.sub(r"\D", "", ref)), re.sub(r"\d", "", ref))
            exp = raw.strip().lstrip(":").strip() if isinstance(raw, str) else raw
            res.cells += 1
            if not same(exp, rep[field], "TEXT"):
                res.diff(ref, field, raw, rep[field])
        res.cells += 1
        if rep["stock_date"] != iso(rep["stock_date_text"] or ""):
            res.diff(s["meta"]["stock_date_text"][0], "stock_date", rep["stock_date_text"], rep["stock_date"])
    out.append(res)

    # daftar kode SAP
    m = s["material"]
    res = Result("WPA REPORT · daftar kode SAP (BA-BB)")
    (code_col, code_h), (desc_col, desc_h) = m["code_col"], m["desc_col"]
    check_header(res, sh, m["header_row"], code_col, code_h)
    check_header(res, sh, m["header_row"], desc_col, desc_h)
    excel, notes = {}, 0  # excel: {baris: (kode, deskripsi)}
    for r in range(m["data_start"], sh.max_row + 1):
        code = sh.get(r, code_col)
        if isinstance(code, (int, float)) and not isinstance(code, bool) and float(code).is_integer():
            excel[r] = (int(code), sh.get(r, desc_col))
        elif not blank(code) or not blank(sh.get(r, desc_col)):
            notes += 1
    # material berisi kode unik (lintas file), jadi dicocokkan per kode, bukan per baris
    db_all = dict(con.execute("SELECT sap_code, description FROM material"))
    db_this = {c for (c,) in con.execute("SELECT sap_code FROM material WHERE file_id=?", (file_id,))}
    excel_codes = {code for code, _ in excel.values()}
    res.excel_rows, res.db_rows = len(excel_codes), len(excel_codes & set(db_all))
    for r, (code, desc) in excel.items():
        res.cells += 2
        if code not in db_all:
            res.diff(f"{code_col}{r}", "kode SAP tidak ada di DB", code, None)
        elif not same(desc, db_all[code], "TEXT"):
            res.diff(f"{desc_col}{r}", f"description kode {code}", desc, db_all[code])
    for code in sorted(db_this - excel_codes):
        res.diff(f"sap_code {code}", "ada di DB tapi tidak ada di daftar Excel")
    dup = len(excel) - len(excel_codes)
    if dup:
        res.notes.append(f"{dup} baris kode duplikat identik (disimpan sekali)")
    if notes:
        res.notes.append(f"{notes} baris catatan (bukan kode) di BA-BB tidak di-ingest")
    out.append(res)

    # lot
    res = Result("WPA REPORT · lot (wpa_stock_lot + quality)")
    for col, (h, _, _) in s["lot_cols"].items():
        check_header(res, sh, s["header_row"], col, h)
    for col, (h, _) in s["quality_cols"].items():
        check_header(res, sh, s["header_row"], col, h)
    first, last = (column_index_from_string(c) for c in s["scan_cols"])
    scan = [get_column_letter(i) for i in range(first, last + 1)]
    excel = [r for r in range(s["data_start"], sh.max_row + 1) if any(not blank(sh.get(r, c)) for c in scan)]
    db = rows_by_source(con, "SELECT * FROM wpa_stock_lot WHERE file_id=?", (file_id,))
    quality = {}
    for lot_id, param, val, col in con.execute(
            "SELECT q.lot_id, q.parameter_code, q.value, q.source_col FROM wpa_stock_lot_quality q "
            "JOIN wpa_stock_lot l USING (lot_id) WHERE l.file_id=?", (file_id,)):
        quality.setdefault(lot_id, {})[col] = (param, val)
    lot_cols = {c: (f, t) for c, (_, f, t) in s["lot_cols"].items()}
    q_cols = {c: p for c, (_, p) in s["quality_cols"].items()}
    for r in compare_rows(res, {r: 1 for r in excel}, db):
        check_values(res, sh, r, db[r], lot_cols)
        check_long(res, sh, r, quality.pop(db[r]["lot_id"], {}), q_cols)
    for lot_id in quality:
        res.diff(f"lot_id {lot_id}", "nilai kualitas tanpa baris lot")
    check_coverage(res, sh, s["data_start"], set(scan) | {code_col, desc_col})
    out.append(res)

    # sheet detail per truk/kontainer
    tests_db = {}
    for rid, param, val, col in con.execute(
            "SELECT t.receipt_id, t.parameter_code, t.value, t.source_col FROM rm_test_result t "
            "JOIN rm_receipt r USING (receipt_id) WHERE r.file_id=?", (file_id,)):
        tests_db.setdefault(rid, {})[col] = (param, val)
    for name, rs in spec.RM_SHEETS.items():
        if name not in wb.sheetnames:
            continue
        sh = Sheet(wb[name])
        res = Result(f"{name} · rm_receipt + rm_test_result")
        ref, h = rs["group_header"]
        check_header(res, sh, 1, ref[0], h)
        for col, (h, hr, _, _) in rs["base_cols"].items():
            check_header(res, sh, hr, col, h)
        for col, (h, _) in rs["test_cols"].items():
            check_header(res, sh, 2, col, h)
        cols = list(rs["base_cols"]) + list(rs["test_cols"])
        excel, block, gap = {}, 1, False
        for r in range(rs["data_start"], sh.max_row + 1):
            if all(blank(sh.get(r, c)) for c in cols):
                gap = gap or bool(excel)
                continue
            if gap:
                block, gap = block + 1, False
            excel[r] = block
        db = rows_by_source(con, "SELECT * FROM rm_receipt WHERE file_id=? AND sheet_name=?", (file_id, name))
        base = {c: (f, t) for c, (_, _, f, t) in rs["base_cols"].items()}
        tcols = {c: p for c, (_, p) in rs["test_cols"].items()}
        for r in compare_rows(res, excel, db):
            rec = db[r]
            res.cells += 2
            if rec["block_no"] != excel[r]:
                res.diff(f"baris {r}", "block_no", excel[r], rec["block_no"])
            if rec["transport_type"] != rs["transport_type"]:
                res.diff(f"baris {r}", "transport_type", rs["transport_type"], rec["transport_type"])
            if "N" not in rs["base_cols"] and rec["remark"] is not None:
                res.diff(f"baris {r}", "remark terisi padahal sheet tidak punya kolom Remark", None, rec["remark"])
            check_values(res, sh, r, rec, base)
            check_long(res, sh, r, tests_db.pop(rec["receipt_id"], {}), tcols)
        blocks = sorted(set(excel.values()))
        res.notes.append(f"{len(blocks)} blok data")
        check_coverage(res, sh, rs["data_start"], set(cols))
        out.append(res)
    for rid in tests_db:
        out[-1].diff(f"receipt_id {rid}", "hasil uji tanpa baris rm_receipt")
    return out


# ---------------------------------------------------------------------- Feed Performance
def validate_feed(con, wb, file_id):
    out = []
    for ws in spec.WIDE_SHEETS:
        sh = Sheet(wb[ws["sheet"]])
        res = Result(f"{ws['sheet']} · {ws['table']}")
        for col, (h, _, _) in ws["cols"].items():
            check_header(res, sh, ws["header_row"], col, h)
        excel = {r: 1 for r in range(ws["data_start"], sh.max_row + 1)
                 if any(not blank(sh.get(r, c)) for c in ws["cols"])}
        db = rows_by_source(con, f"SELECT * FROM {ws['table']} WHERE file_id=?", (file_id,))
        cols = {c: (f, t) for c, (_, f, t) in ws["cols"].items()}
        for r in compare_rows(res, excel, db):
            check_values(res, sh, r, db[r], cols)
        check_coverage(res, sh, ws["data_start"], set(ws["cols"]), ws["skip_cols"])
        out.append(res)
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--input", default=str(ROOT / "docs_revised"))
    ap.add_argument("--db", default=str(HERE / "aft.sqlite"))
    args = ap.parse_args()

    con = sqlite3.connect(args.db)
    failed = False
    for path in sorted(p for p in Path(args.input).glob("*.xlsx") if not p.name.startswith("~$")):
        print(f"\n=== {path.name}")
        row = con.execute("SELECT file_id, file_kind FROM source_file WHERE file_name=?", (path.name,)).fetchone()
        if row is None:
            print("  GAGAL: file tidak ada di source_file")
            failed = True
            continue
        file_id, kind = row
        wb = load_workbook(path, read_only=True, data_only=True)
        handled = ({spec.WPA_REPORT["sheet"], *spec.RM_SHEETS} if kind == "WPA"
                   else {ws["sheet"] for ws in spec.WIDE_SHEETS})
        unknown = [n for n in wb.sheetnames if n not in handled | spec.SKIPPED_SHEETS]
        skipped = [n for n in wb.sheetnames if n in spec.SKIPPED_SHEETS]
        results = validate_wpa(con, wb, file_id) if kind == "WPA" else validate_feed(con, wb, file_id)
        wb.close()

        print(f"  {'bagian':45s} {'baris Excel':>11s} {'baris DB':>9s} {'sel dicek':>10s} {'beda':>5s} {'tak-mapping':>12s}")
        for res in results:
            print(f"  {res.label:45s} {res.excel_rows:11d} {res.db_rows:9d} {res.cells:10d} "
                  f"{len(res.diffs):5d} {len(res.unmapped):12d}  {'OK' if res.ok else 'GAGAL'}")
            for d in res.diffs[:MAX_SHOW]:
                print(f"      beda: {d}")
            for u in res.unmapped[:MAX_SHOW]:
                print(f"      tak-mapping: {u}")
            for n in res.notes[:3]:
                print(f"      catatan: {n}")
            failed |= not res.ok
        if skipped:
            print(f"  sheet tidak di-ingest (rekap/pivot): {', '.join(skipped)}")
        if unknown:
            print(f"  GAGAL: sheet tidak dikenal: {', '.join(unknown)}")
            failed = True
        issues = list(con.execute("SELECT sheet_name, cell_ref, raw_value, reason FROM ingest_issue WHERE file_id=?", (file_id,)))
        for sheet, ref, raw, reason in issues[:MAX_SHOW]:
            print(f"  ingest_issue: {sheet}!{ref} {raw!r} -> {reason}")
        failed |= bool(issues)

    print("\nHASIL:", "GAGAL" if failed else "SEMUA COCOK")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
