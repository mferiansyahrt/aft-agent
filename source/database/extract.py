"""Ekstrak record dari Excel CP Prima sesuai sheet_spec.py.

Excel dibaca dengan nilai hasil rumus (data_only=True). Setiap record menyimpan
asal selnya (sheet, baris, kolom) supaya bisa dicek balik oleh validate.py.
"""
from datetime import date, datetime

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

import sheet_spec as spec

DATE_FORMATS = ("%d %B %Y", "%d %b %Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y")


def is_empty(v):
    return v is None or (isinstance(v, str) and not v.strip())


def _to_iso(v):
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == datetime.min.time() else v.isoformat(sep=" ")
    return v.isoformat()


def convert(v, typ):
    """Ubah nilai sel ke tipe kolom. Hasil: (nilai, alasan_gagal atau None)."""
    if is_empty(v):
        return None, None
    if typ == "TEXT":
        if isinstance(v, str):
            return v.strip(), None
        if isinstance(v, (datetime, date)):
            return _to_iso(v), None
        if isinstance(v, float) and v.is_integer():
            return str(int(v)), None  # 111223135.0 -> "111223135"
        return str(v), None
    if typ in ("REAL", "INTEGER"):
        if isinstance(v, bool) or not isinstance(v, (int, float, str)):
            return None, f"tipe {type(v).__name__} bukan angka"
        try:
            x = float(v.strip()) if isinstance(v, str) else v
        except ValueError:
            return None, "bukan angka"
        if typ == "REAL":
            return float(x), None
        if float(x).is_integer():
            return int(x), None
        return None, "bukan bilangan bulat"
    if typ == "DATE":
        if isinstance(v, (datetime, date)):
            return _to_iso(v), None
        if isinstance(v, str):
            for fmt in DATE_FORMATS:
                try:
                    return datetime.strptime(v.strip(), fmt).date().isoformat(), None
                except ValueError:
                    pass
            return None, "format tanggal tidak dikenal"
        return None, f"tipe {type(v).__name__} bukan tanggal"
    raise ValueError(f"tipe tidak dikenal: {typ}")


class Sheet:
    """Isi sheet sebagai {nomor_baris: tuple nilai}, dibaca sekali dengan iter_rows."""

    def __init__(self, ws):
        self.name = ws.title
        self.rows = {i: r for i, r in enumerate(ws.iter_rows(values_only=True), start=1)}
        self.max_row = max(self.rows, default=0)

    def get(self, row, col):
        r = self.rows.get(row)
        i = column_index_from_string(col) - 1
        return r[i] if r is not None and i < len(r) else None

    def get_ref(self, ref):
        col = "".join(ch for ch in ref if ch.isalpha())
        return self.get(int(ref[len(col):]), col)


class Extractor:
    def __init__(self, path):
        self.path = path
        self.wb = load_workbook(path, read_only=True, data_only=True)
        self.kind = spec.file_kind(self.wb.sheetnames)
        self.issues = []

    def _val(self, sh, row, col, typ):
        raw = sh.get(row, col)
        val, err = convert(raw, typ)
        if err:
            self.issues.append({"sheet_name": sh.name, "cell_ref": f"{col}{row}",
                                "raw_value": str(raw), "expected_type": typ, "reason": err})
        return val

    def run(self):
        if self.kind == "WPA":
            out = self._wpa()
        elif self.kind == "FEED_PERF":
            out = self._feed()
        else:
            raise ValueError(f"jenis file tidak dikenal: {self.path}")
        out["kind"] = self.kind
        out["issues"] = self.issues
        self.wb.close()
        return out

    # ------------------------------------------------------------------ WPA
    def _wpa(self):
        s = spec.WPA_REPORT
        sh = Sheet(self.wb[s["sheet"]])

        report = {}
        for field, (ref, typ) in s["meta"].items():
            raw = sh.get_ref(ref)
            if isinstance(raw, str) and raw.strip().startswith(":"):
                raw = raw.strip()[1:]
            report[field] = convert(raw, typ)[0]
        report["stock_date"], err = convert(report["stock_date_text"], "DATE")
        if err:
            self.issues.append({"sheet_name": sh.name, "cell_ref": s["meta"]["stock_date_text"][0],
                                "raw_value": report["stock_date_text"], "expected_type": "DATE", "reason": err})

        m = s["material"]
        code_col, desc_col = m["code_col"][0], m["desc_col"][0]
        materials = []
        for r in range(m["data_start"], sh.max_row + 1):
            code = sh.get(r, code_col)
            if isinstance(code, (int, float)) and not isinstance(code, bool) and float(code).is_integer():
                materials.append({"sap_code": int(code),
                                  "description": convert(sh.get(r, desc_col), "TEXT")[0],
                                  "source_row": r})

        first, last = (column_index_from_string(c) for c in s["scan_cols"])
        lots = []
        for r in range(s["data_start"], sh.max_row + 1):
            row = sh.rows.get(r) or ()
            if all(is_empty(v) for v in row[first - 1:last]):
                continue
            lot = {"source_row": r}
            for col, (_, field, typ) in s["lot_cols"].items():
                lot[field] = self._val(sh, r, col, typ)
            lot["quality"] = []
            for col, (_, param) in s["quality_cols"].items():
                v = self._val(sh, r, col, "REAL")
                if v is not None:
                    lot["quality"].append((param, v, col))
            lots.append(lot)

        receipts = []
        for name, rs in spec.RM_SHEETS.items():
            if name not in self.wb.sheetnames:
                continue
            receipts.extend(self._rm_sheet(Sheet(self.wb[name]), rs))

        return {"report": report, "materials": materials, "lots": lots, "receipts": receipts}

    def _rm_sheet(self, sh, rs):
        cols = list(rs["base_cols"]) + list(rs["test_cols"])
        out, block, gap = [], 1, False
        for r in range(rs["data_start"], sh.max_row + 1):
            if all(is_empty(sh.get(r, c)) for c in cols):
                gap = bool(out)
                continue
            if gap:
                block += 1
                gap = False
            rec = {"sheet_name": sh.name, "source_row": r, "block_no": block,
                   "transport_type": rs["transport_type"], "remark": None}
            for col, (_, _, field, typ) in rs["base_cols"].items():
                rec[field] = self._val(sh, r, col, typ)
            rec["tests"] = []
            for col, (_, param) in rs["test_cols"].items():
                v = self._val(sh, r, col, "REAL")
                if v is not None:
                    rec["tests"].append((param, v, col))
            out.append(rec)
        return out

    # ------------------------------------------------------------------ Feed Performance
    def _feed(self):
        out = {}
        for ws_spec in spec.WIDE_SHEETS:
            sh = Sheet(self.wb[ws_spec["sheet"]])
            rows = []
            for r in range(ws_spec["data_start"], sh.max_row + 1):
                if all(is_empty(sh.get(r, c)) for c in ws_spec["cols"]):
                    continue
                rec = {"source_row": r}
                for col, (_, field, typ) in ws_spec["cols"].items():
                    rec[field] = self._val(sh, r, col, typ)
                rows.append(rec)
            out[ws_spec["table"]] = rows
        return out


def extract(path):
    return Extractor(path).run()
